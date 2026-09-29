"""Builds a Word evidence document from the agent's execution log."""
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.shared import Inches

REPORTS_DIR = Path(__file__).resolve().parent.parent / "reports"
REPORTS_DIR.mkdir(exist_ok=True)


def generate_report(scenario: str, execution_log: list[dict]) -> str:
    doc = Document()
    doc.add_heading("QA Automation Execution Report", level=1)
    doc.add_paragraph(f"Scenario: {scenario}")
    doc.add_paragraph(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")

    passed = sum(1 for entry in execution_log if entry["status"] == "pass")
    failed = sum(1 for entry in execution_log if entry["status"] == "fail")
    doc.add_paragraph(f"Total Steps: {len(execution_log)} | Passed: {passed} | Failed: {failed}")

    for entry in execution_log:
        doc.add_heading(f"Step: {entry['step']} - {entry['status'].upper()}", level=2)
        doc.add_paragraph(f"Time: {entry['timestamp']}")
        doc.add_paragraph(f"Details: {entry['message']}")
        if entry.get("screenshot") and Path(entry["screenshot"]).exists():
            doc.add_picture(entry["screenshot"], width=Inches(5.5))

    filename = REPORTS_DIR / f"execution_report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx"
    doc.save(str(filename))
    return str(filename)
